import requests
import time

from docx import Document as DocxDoc

doc = DocxDoc()
doc.add_heading("DocuMind Test Document", 0)
doc.add_paragraph("This document is used to test the DocuMind ingestion pipeline.")
doc.add_paragraph("It covers topics like artificial intelligence, retrieval augmented generation, and vector search.")
doc.add_paragraph("The system should extract this text, chunk it, embed it, and index it into Qdrant.")
doc.save("test_document.docx")
print("Created test_document.docx")

with open("test_document.docx", "rb") as f:
    r = requests.post(
        "http://localhost:8000/api/v1/chats/3/documents",
        files={"file": ("test_document.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
print("Upload:", r.status_code, r.json())
doc_id = r.json()["id"]

for i in range(15):
    time.sleep(2)
    r2 = requests.get("http://localhost:8000/api/v1/chats/3/documents")
    for d in r2.json():
        if d["id"] == doc_id:
            status = d["status"]
            error = d.get("error_message") or ""
            print(f"  [{i*2}s] {status} {error}")
            if status in ("READY", "FAILED"):
                raise SystemExit(0 if status == "READY" else 1)
